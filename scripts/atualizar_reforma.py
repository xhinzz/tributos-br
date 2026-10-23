"""cClassTrib (IT 2025.002, Portal da NF-e) e NCMs dos anexos da LC 214/2025 (Planalto).

    python scripts/atualizar_reforma.py [cclass.xlsx] [lcp214.htm]

Requer openpyxl. Saída: tributos_br/dados/cclass_trib.csv e dados/listas/lc214_anexo_<n>.csv.
"""

import csv
import io
import re
import sys
import urllib.request
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from atualizar_convenios import _Tabelas, _vigentes

CCLASS = "https://www.nfe.fazenda.gov.br/portal/exibirArquivo.aspx?conteudo=Cd5xq2ZVuEg="
LC214 = "https://www.planalto.gov.br/ccivil_03/leis/lcp/lcp214.htm"
DADOS = Path(__file__).resolve().parent.parent / "tributos_br" / "dados"
ROMANOS = {
    r: i
    for i, r in enumerate(
        ("I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII", "XIV", "XV"), 1
    )
}
# II, III e X são serviços (NBS); o XIV foi revogado pela LC 227/2026
ANEXOS_DE_BENS = {1, 4, 5, 6, 7, 8, 9, 11, 12, 13, 15}
# a lei escreve 0306.1, 07.01, 07.02.00.00 e "capítulo 10"; NBS (1.2201.1) tem ponto antes e fica de fora
NCM = re.compile(r"(?<![\d.])(\d{4}|\d{2}\.\d{2})((?:\.\d{1,2}){0,2})(?!\d)")
CAPITULO = re.compile(r"cap[íi]tulos? (\d{1,2})\b", re.I)


def _codigos(texto: str) -> list[str]:
    achados = []
    for inicio, resto in NCM.findall(texto):
        if not resto and "." not in inicio and 1950 <= int(inicio) <= 2049:
            continue  # ano solto ("Lei nº 10.711, de 2003")
        achados.append((inicio + resto).replace(".", ""))
    return achados + [c.zfill(2) for c in CAPITULO.findall(texto)]


def _ncms(texto: str) -> list[str]:
    """NCMs do item; os citados em "exceto ..." saem com "!" na frente."""
    excecoes = [f"!{c}" for trecho in re.findall(r"exceto[^;)]*", texto) for c in _codigos(trecho)]
    return list(dict.fromkeys(_codigos(re.sub(r"exceto[^;)]*", "", texto)) + excecoes))


def _baixar(url: str) -> bytes:
    pedido = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(pedido, timeout=120) as resposta:
        return resposta.read()


def _data(valor) -> str:
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    return str(valor or "")[:10]


def cclass(conteudo: bytes) -> list[list[str]]:
    planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).worksheets[0]
    linhas = planilha.iter_rows(values_only=True)
    cabecalho = {nome: i for i, nome in enumerate(next(linhas)) if nome}
    saida = []
    for linha in linhas:
        if not linha[cabecalho["cClassTrib"]]:
            continue
        valor = dict(zip(cabecalho, (linha[i] for i in cabecalho.values()), strict=True)).get
        anexo = str(valor("ANEXO") or "")
        saida.append(
            [
                str(valor("CST-IBS/CBS")).zfill(3),
                str(valor("cClassTrib")).zfill(6),
                " ".join(str(valor("Nome cClassTrib")).split()),
                valor("Tipo de Alíquota"),
                str(valor("pRedIBS") or 0),
                str(valor("pRedCBS") or 0),
                str(valor("indNFe") or 0),
                # só os anexos da LC 214 (1 a 15); o resto da coluna são códigos de NBS
                anexo if anexo.isdigit() and int(anexo) in ANEXOS_DE_BENS else "",
                _data(valor("dIniVig")),
                _data(valor("dFimVig")),
                str(valor("LC 214/25") or ""),
            ]
        )
    return saida


def anexos(html: str) -> dict[int, list[list[str]]]:
    p = _Tabelas()
    p.feed(html)
    saida = {}
    for tabela in p.tabelas:
        # o título do Anexo I cita o XV entre parênteses: vale só "ANEXO N" seguido de palavra em caixa alta
        achado = re.findall(r"ANEXO ([IVX]+) (?=[A-ZÀ-Ú]{3,})", tabela["titulo"])
        numero = ROMANOS.get(achado[-1]) if achado else None
        if numero not in ANEXOS_DE_BENS or numero in saida or tabela["linhas"][0][0] != "ITEM":
            continue
        itens = []
        for linha in _vigentes(tabela["linhas"][1:]):
            itens += [[linha[0], ncm, linha[1][:200]] for ncm in _ncms(" ".join(linha[1:]))]
        if itens:
            saida[numero] = itens
    return saida


def _gravar(caminho: Path, cabecalho, linhas):
    caminho.parent.mkdir(parents=True, exist_ok=True)
    with caminho.open("w", encoding="utf-8", newline="") as f:
        escritor = csv.writer(f, lineterminator="\n")
        escritor.writerow(cabecalho)
        escritor.writerows(linhas)
    print(f"{caminho.relative_to(DADOS)}: {len(linhas)} linhas")


def main():
    planilha = Path(sys.argv[1]).read_bytes() if len(sys.argv) > 1 else _baixar(CCLASS)
    html = (Path(sys.argv[2]).read_bytes() if len(sys.argv) > 2 else _baixar(LC214)).decode("cp1252")
    _gravar(
        DADOS / "cclass_trib.csv",
        [
            "cst",
            "cclass_trib",
            "nome",
            "tipo",
            "p_red_ibs",
            "p_red_cbs",
            "nfe",
            "anexo",
            "vigencia_inicio",
            "vigencia_fim",
            "fundamento",
        ],
        cclass(planilha),
    )
    for numero, itens in sorted(anexos(html).items()):
        _gravar(DADOS / "listas" / f"lc214_anexo_{numero}.csv", ["item", "ncm", "descricao"], itens)


if __name__ == "__main__":
    main()
