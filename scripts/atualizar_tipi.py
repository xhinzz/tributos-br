"""Baixa a TIPI oficial da Receita Federal (xlsx) e gera tributos_br/dados/tipi.csv.

Uso: python scripts/atualizar_tipi.py [caminho.xlsx]   (requer openpyxl)

A planilha mistura tipos na coluna de alíquota (número com ruído de float, 'NT', vazio nas linhas de
agrupamento); só ficam as linhas de NCM com 8 dígitos.
"""

import csv
import io
import sys
import urllib.request
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

URL = "https://www.gov.br/receitafederal/pt-br/acesso-a-informacao/legislacao/documentos-e-arquivos/tipi.xlsx"
DESTINO = Path(__file__).resolve().parent.parent / "tributos_br" / "dados" / "tipi.csv"


def baixar() -> bytes:
    # o gov.br responde 403 sem User-Agent de navegador
    pedido = urllib.request.Request(URL, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"})
    with urllib.request.urlopen(pedido, timeout=60) as resposta:
        return resposta.read()


def _texto(valor: Decimal) -> str:
    return format(valor.normalize(), "f")


def linhas(conteudo: bytes):
    planilha = load_workbook(io.BytesIO(conteudo), read_only=True, data_only=True).worksheets[0]
    for ncm, ex, _descricao, aliquota, *_ in planilha.iter_rows(min_row=9, values_only=True):
        ncm = str(ncm or "").replace(".", "").strip()
        if len(ncm) != 8 or not ncm.isdigit() or aliquota in (None, ""):
            continue
        if isinstance(aliquota, str):
            aliquota = aliquota.strip().upper()
            if aliquota != "NT":
                aliquota = _texto(Decimal(aliquota.replace(",", ".")))
        else:
            aliquota = _texto(round(Decimal(repr(aliquota)), 2))
        yield ncm, str(ex or "").strip(), aliquota


def main():
    conteudo = Path(sys.argv[1]).read_bytes() if len(sys.argv) > 1 else baixar()
    vistas = {}
    for ncm, ex, aliquota in linhas(conteudo):
        vistas[(ncm, ex)] = aliquota
    with DESTINO.open("w", encoding="utf-8", newline="") as f:
        escritor = csv.writer(f, lineterminator="\n")
        escritor.writerow(["ncm", "ex", "aliquota"])
        escritor.writerows((ncm, ex, a) for (ncm, ex), a in sorted(vistas.items()))
    print(f"{DESTINO.name}: {len(vistas)} linhas")


if __name__ == "__main__":
    main()
